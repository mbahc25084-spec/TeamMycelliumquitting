from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Mycelium_Final_Handoff.docx"

PLUM = "2F114F"
PLUM_DARK = "130621"
ORCHID = "8B44D9"
GOLD = "C9861A"
CYAN = "0E7490"
MINT = "16865C"
ROSE = "B4235A"
PALE = "F7F3FA"
LIGHT = "EDE6F3"
GRID = "D9D9D9"
BLACK = "000000"
GRAY = "53515A"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def cell_border(cell, color=GRID, size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=95, start=110, bottom=95, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + side))
        if node is None:
            node = OxmlElement("w:" + side)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    p_pr.append(keep)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_run(paragraph, text, *, bold=False, size=None, color=None, italic=False):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    return run


def add_body(doc, text, *, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.18
    add_run(p, text, size=10.7, color=BLACK)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(15 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = add_run(p, text, bold=True, size=15 if level == 1 else 11.5, color=BLACK)
    run.font.name = "Aptos Display" if level == 1 else "Aptos"
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.1
    add_run(p, text, size=10.3, color=BLACK)
    return p


def add_table(doc, headers, rows, widths, font_size=8.9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, label in enumerate(headers):
        cell = header.cells[idx]
        cell.width = Inches(widths[idx])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        shade(cell, PLUM)
        cell_border(cell, PLUM)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        add_run(p, label, bold=True, size=font_size, color="FFFFFF")
    for r_idx, values in enumerate(rows):
        row = table.add_row()
        for idx, value in enumerate(values):
            cell = row.cells[idx]
            cell.width = Inches(widths[idx])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            shade(cell, "FFFFFF" if r_idx % 2 == 0 else PALE)
            cell_border(cell)
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.06
            add_run(p, str(value), size=font_size, color=BLACK)
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_together = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(6)
    add_run(p, "Mycelium prototype handoff  |  Offline demo", size=8, color=GRAY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    p._p.append(fld_char1)
    p._p.append(instr)
    p._p.append(fld_char2)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    footer(section)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    styles["Normal"].font.size = Pt(10.7)
    for style_name in ("Heading 1", "Heading 2", "Title"):
        style = styles[style_name]
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")

    # Cover
    accent = doc.add_table(rows=1, cols=3)
    accent.autofit = False
    for cell, fill in zip(accent.rows[0].cells, (PLUM, ORCHID, GOLD)):
        cell.width = Inches(2.2)
        shade(cell, fill)
        cell_border(cell, fill, "0")
        set_cell_margins(cell, 25, 10, 25, 10)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_before = Pt(26)
    title.paragraph_format.space_after = Pt(8)
    add_run(title, "Mycelium Prototype Final Handoff", bold=True, size=29, color=BLACK)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(20)
    add_run(subtitle, "Offline Authority Aware Booking Prototype", size=15, color=BLACK)
    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(15)
    intro.paragraph_format.line_spacing = 1.22
    add_run(intro, "This handoff records the final Mycelium prototype and its supplied source pack. The prototype watches simulated booking opportunities, completes reversible preparation, and pauses when a person owns the next decision. It is built for a deterministic offline demo with visible provenance and an append only audit trace.", size=12, color=BLACK)

    cover_table = doc.add_table(rows=4, cols=2)
    cover_table.autofit = False
    cover_rows = [
        ("Delivery", "Static browser prototype and final handoff"),
        ("Decision model", "Authority graph consent envelope and policy trace"),
        ("Data posture", "LOCAL and SIMULATED labels visible throughout"),
        ("Primary review path", "Scenario S2 one dimension outside the envelope"),
    ]
    for i, (label, value) in enumerate(cover_rows):
        cells = cover_table.rows[i].cells
        for cell, fill in zip(cells, (LIGHT, "FFFFFF")):
            cell.width = Inches(1.8 if cell is cells[0] else 4.8)
            shade(cell, fill)
            cell_border(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, 115, 125, 115, 125)
        p0 = cells[0].paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        add_run(p0, label, bold=True, size=10.1, color=PLUM_DARK)
        p1 = cells[1].paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        add_run(p1, value, size=10.1, color=BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(14)
    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(8)
    note.paragraph_format.space_after = Pt(0)
    add_run(note, "Scope note  ", bold=True, size=9.5, color=PLUM)
    add_run(note, "The appended source pack is reference material. Its fixture data is not presented as live inventory, payments, research, or production performance evidence.", size=9.5, color=GRAY)

    doc.add_page_break()
    add_heading(doc, "Prototype Summary")
    add_body(doc, "Mycelium gives an autonomous booking assistant a bounded role. It can search, compare, watch, and prepare a transaction, but it does not treat a person’s preferences, identity, money, or irreversible commitment as interchangeable authority. The active decision is computed from scenario fixtures and the current append only event log.")
    add_heading(doc, "What A Reviewer Can Test", 2)
    for bullet in [
        "Select any deterministic scenario from S1 through S10, including an inside envelope confirmation, cross owner money lock, budget refusal, healthcare scheduling case, revocation, and ambiguous voice reply.",
        "Open Authority Lab and change ownership, consent, identity mapping, reachability, fares, times, and budget. The visible verdict and rule trace recompute from the local policy.",
        "Inspect Agent Brain and scrub the event sequence. The same replayed events produce the current decision object and audit timeline.",
        "Complete an approval flow to reach the pre capture guards and Decision Receipt. Payment and inventory are clearly labelled simulated.",
    ]:
        add_bullet(doc, bullet)
    add_heading(doc, "Decision Lifecycle", 2)
    add_table(doc, ["Step", "Prototype behaviour", "Evidence surface"], [
        ["Detect", "A simulated opportunity is found during a watch.", "Watch console and event trace"],
        ["Evaluate", "The policy resolves the domain owner and checks bounded consent.", "Rule trace and decision object"],
        ["Prepare", "The prototype stages reversible work and a free temporary hold.", "Prepared transaction state"],
        ["Pause or proceed", "A qualifying envelope can execute. An unresolved or cross owner decision asks or locks.", "Approver pane and authority lock"],
        ["Guard", "Identity and envelope validity are checked before capture.", "Pre capture guards and audit events"],
        ["Record", "The receipt names the decision owner and the final state.", "Decision Receipt"],
    ], [0.74, 3.57, 2.06])

    add_heading(doc, "Scenario Coverage")
    add_body(doc, "The ten scenario fixtures are review cases, not claims about real users or services. They expose different authority and safety decisions through one consistent interaction model.")
    scenarios = [
        ["S1", "Inside the envelope", "Confirm under recorded consent"],
        ["S2", "One dimension outside", "Ask Dad about the 22 20 departure"],
        ["S3", "Approver unreachable", "Release using the declared fallback"],
        ["S4", "Slot expires", "Release the free hold and keep watching"],
        ["S5", "Money decision", "Lock because the booker is not the money owner"],
        ["S6", "Budget breach", "Refuse above the hard ceiling"],
        ["S7", "Two qualifying slots", "Ask rather than choose a tie"],
        ["S8", "Healthcare scheduling", "Ask Amma about her appointment time"],
        ["S9", "Revoked mid flight", "Void the mandate before capture"],
        ["S10", "Ambiguous voice reply", "Re ask one decision dimension"],
    ]
    add_table(doc, ["ID", "Fixture", "Expected policy result"], scenarios, [0.55, 2.05, 3.77], font_size=8.7)

    add_heading(doc, "Policy And Audit Integrity")
    add_body(doc, "The supplied engine source defines a policy family from R1 through R15. The final prototype exposes its corresponding authority checks through the live rule trace. It keeps a separate distinction between reversible work, bounded consent, and money or commitment decisions that require the correct owner.")
    add_table(doc, ["Control", "How it appears in the prototype", "Provenance"], [
        ["Decision authority", "Seven decision domains identify an owner for discovery, selection, schedule, identity, money, commitment, and disclosure.", "LOCAL"],
        ["Consent envelope", "A recorded boundary checks fare, class, departure time, and cancellation against the fixture opportunity.", "LOCAL"],
        ["Opportunity and payment", "Rail inventory, holds, mandate states, payment capture, and voice delivery are demo fixtures.", "SIMULATED"],
        ["Audit trace", "User visible state changes append an event before the interface re renders.", "LOCAL"],
        ["Replay", "The trace can be scrubbed to recompute a prior point in the same event sequence.", "LOCAL"],
    ], [1.27, 4.15, 0.95])
    add_heading(doc, "Visible Constraints", 2)
    add_body(doc, "The delivery deliberately makes its limits visible. It has no live API dependency, it does not claim production inventory or payment execution, and it does not present fixture values as research or business results. The interface uses LOCAL and SIMULATED badges so a reviewer can distinguish the policy demonstration from a live service.")
    add_heading(doc, "Static Prototype Adaptation", 2)
    add_body(doc, "The source pack describes a React and TypeScript build plan. This delivery is a self contained static HTML, CSS, and JavaScript prototype so it can run offline without installing a toolchain or contacting a service. The source pack is retained in the final delivery as reference material; the runtime behaviour is implemented in the static application files.")

    add_heading(doc, "Visual Direction")
    add_body(doc, "Mycelium uses a light booker surface above and a dark authority surface below. Colour carries decision meaning: violet identifies authority, ochre identifies money, cyan identifies time, rose identifies refusal, and mint identifies confirmed or safe states. The opening and authority lock moments use inline SVG characters and props so the visual layer remains inside the offline bundle.")
    add_table(doc, ["Semantic role", "Visual treatment", "Used for"], [
        ["Authority", "Violet plum", "Owners and consent boundaries"],
        ["Money", "Ocher gold", "Payment rail and money owner"],
        ["Time", "Cyan", "Hold countdown and timing constraints"],
        ["Refusal", "Rose", "Budget refusal and invalid actions"],
        ["Confirmation", "Mint", "Confirmed actions and passing guards"],
    ], [1.25, 1.75, 3.37])

    add_heading(doc, "Demo Guide")
    add_body(doc, "Start with S1 to show that the assistant can complete a fully bounded booking without an approval request. Then switch to S2. It contains a fare, class, and cancellation condition that pass, while the 22 20 departure falls outside the approved window. The app prepares the transaction but leaves money with Dad and asks only about the unresolved time.")
    add_heading(doc, "Suggested Review Sequence", 2)
    add_table(doc, ["Moment", "Action", "What the reviewer should observe"], [
        ["1", "Choose S1", "The consent envelope covers the fixture and confirmation can continue without interrupting Dad."],
        ["2", "Choose S2", "One timing condition fails. The interface shows the narrow question rather than a broad booking form."],
        ["3", "Open Authority Lab", "Give the booker money authority, then restore it. The verdict changes with the selected authority."],
        ["4", "Run revocation attack", "The mandate voids before capture because validity is checked at the final step."],
        ["5", "Open Agent Brain", "The append only event list and replay control expose the same decision history used by the interface."],
        ["6", "Open Decision Receipt", "The receipt states who decided the critical steps and whether the booking fixture was confirmed."],
    ], [0.5, 1.55, 4.32])
    add_heading(doc, "Controls", 2)
    for bullet in [
        "Scenario picker in the header selects S1 through S10. Keyboard keys 1 through 9 select S1 through S9, and 0 selects S10.",
        "Authority Lab exposes ownership, consent, identity mapping, reachability, envelope values, found values, budget, and five attack presets.",
        "Agent Brain includes an event replay slider. Decision Receipt and Trust Surfaces are available from the primary navigation.",
        "Press P outside an input to toggle projector mode. This change also becomes an audit event.",
    ]:
        add_bullet(doc, bullet)

    add_heading(doc, "Delivery Contents")
    add_table(doc, ["Path", "Purpose"], [
        ["index html", "Static application entry point"],
        ["styles css", "Responsive visual system and semantic colour treatment"],
        ["script js", "Fixtures event log policy decisions scenarios and UI rendering"],
        ["docs ART DIRECTION md", "Brand visual and illustration direction"],
        ["docs source pack", "Supplied engine fixtures and feature reference documents"],
        ["README md", "Local launch and reviewer controls"],
        ["Mycelium Final Handoff docx", "This final review document"],
    ], [2.3, 4.07])
    closing = doc.add_paragraph()
    closing.paragraph_format.space_before = Pt(12)
    closing.paragraph_format.space_after = Pt(0)
    closing.paragraph_format.line_spacing = 1.18
    add_run(closing, "The final package is ready for local review. It demonstrates a specific safety proposition: the agent can move quickly around a booking while preserving the person who owns the decision.", size=10.7, color=BLACK)

    doc.core_properties.title = "Mycelium Prototype Final Handoff"
    doc.core_properties.subject = "Offline authority aware booking prototype"
    doc.core_properties.author = "Mycelium prototype team"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
